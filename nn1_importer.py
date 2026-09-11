"""
nn1_importer.py — 工程部 NN1 Project Excel 匯入 + Master List 比對

用途：上司提供的「2026 NN1 Project.xlsx」→ 比對 quotation_registry → 組會簽表 payload。
"""
from __future__ import annotations

import os
import re
from datetime import datetime

import openpyxl

import database as db
from master_ref import normalize_person_code, person_display_name

REF_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Ref')

SIGNOFF_PM_FIELDS = (
    'submit_date',
    'tender_date',
    'bid_deadline',
    'bid_deadline_ampm',
    'bid_deadline_meridiem',
    'approval_category',
    'approval_other',
    'attachments',
    'remark',
)

SIGNOFF_CONTENT_FIELDS = (
    'quotation_no',
    'project_name',
    'client_name',
    'amount',
    'contract_start_date',
    'contract_end_date',
    'expected_period',
    'contract_months',
    'subcon_name',
    'subcon_quotation',
    'person_in_charge',
    'person_code',
)

MASTER_COMPARE_FIELDS = (
    ('description', '工程名稱', 'description'),
    ('client_name', '供方名稱', 'client_name'),
    ('quoted_amount', '工程金額', 'amount'),
    ('awarded_amount', '中標金額', 'awarded_amount'),
    ('start_date', '開工日期', 'start_date'),
    ('completion_date', '完工日期', 'completion_date'),
    ('person_in_charge', '項目負責人', 'person_in_charge'),
    ('site_name', '屋苑/地點', 'site_name'),
)


def _safe_str(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if not s or s.lower() in ('none', 'nan'):
        return None
    if s.upper() in ('NA', 'N/A', '-', '—', '——'):
        return None
    return s


def _safe_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    s = _safe_str(val)
    if not s:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d/%m/%y', '%Y/%m/%d'):
        try:
            return datetime.strptime(s[:10], fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    if re.match(r'\d{4}-\d{2}-\d{2}', s):
        return s[:10]
    return None


def _safe_float(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if isinstance(val, float) and val != val:
            return None
        return float(val)
    s = _safe_str(val)
    if not s:
        return None
    s = s.replace(',', '').replace('$', '').replace('HK', '')
    try:
        return float(s)
    except ValueError:
        return None


def normalize_nn1_quotation(raw):
    """NN1 編號 → Master 比對鍵。Q.0080/25 → Q080/25；Q.1073/26 → Q1073/26。"""
    s = _safe_str(raw)
    if not s:
        return None
    s = re.sub(r'\s+', '', s)
    s = re.sub(r'^([QTC])\.', r'\1', s, flags=re.I)
    m = re.search(r'([QTC])(\d+)/(\d{2,4})', s, re.I)
    if not m:
        return None
    letter = m.group(1).upper()
    num = int(m.group(2))
    year = m.group(3)
    if len(year) == 4:
        year = year[-2:]
    num_str = str(num).zfill(3) if num < 1000 else str(num)
    return f'{letter}{num_str}/{year}'


def build_quotation_no(raw, person_code=None):
    """組完整 Master 編號 MS/Q1073/26/dc。"""
    core = normalize_nn1_quotation(raw)
    if not core:
        return None
    q = f'MS/{core}'
    pc = normalize_person_code(person_code)
    if pc:
        q = f'{q}/{pc}'
    return q


def _quotation_candidates(raw, person_code=None):
    core = normalize_nn1_quotation(raw)
    if not core:
        return []
    out = [f'MS/{core}']
    pc = normalize_person_code(person_code)
    if pc:
        out.append(f'MS/{core}/{pc}')
    return out


def find_master_record(raw, person_code=None):
    for q in _quotation_candidates(raw, person_code):
        row = db.get_quotation_by_no(q)
        if row:
            return dict(row), q
    core = normalize_nn1_quotation(raw)
    if not core:
        return None, None
    conn = db.get_conn()
    row = conn.execute(
        "SELECT * FROM quotation_registry WHERE quotation_no LIKE ? ORDER BY quotation_no LIMIT 1",
        (f'%/{core}%',),
    ).fetchone()
    conn.close()
    if row:
        d = dict(row)
        return d, d.get('quotation_no')
    return None, None


def _infer_client(project_name, master_row):
    if master_row and master_row.get('client_name'):
        return master_row['client_name']
    name = (project_name or '').upper()
    if 'MTR' in name or '港鐵' in (project_name or ''):
        return '香港鐵路有限公司'
    return None


def _pick_amount(nn1_amount, master_row):
    if master_row:
        for key in ('awarded_amount', 'quoted_amount'):
            v = master_row.get(key)
            if v not in (None, '', 0):
                return float(v)
    if nn1_amount is not None:
        return float(nn1_amount)
    return None


def _months_between(start, end):
    if not start or not end:
        return None
    try:
        s = datetime.strptime(start[:10], '%Y-%m-%d')
        e = datetime.strptime(end[:10], '%Y-%m-%d')
        months = (e.year - s.year) * 12 + (e.month - s.month)
        if e.day >= s.day:
            months += 1
        return max(months, 1)
    except ValueError:
        return None


def _format_period_months(val) -> str | None:
    """合約期（月）顯示：整數去小數點，小數保留必要位。"""
    if val in (None, ''):
        return None
    try:
        n = float(val)
    except (TypeError, ValueError):
        return None
    if n <= 0:
        return None
    if n == int(n):
        return str(int(n))
    return str(n).rstrip('0').rstrip('.')


def _parse_period_months(val):
    """合約期（月）：支援 1、5、0.5 或舊式 '6 months'。"""
    if val in (None, ''):
        return None
    if isinstance(val, (int, float)):
        n = float(val)
        return n if n > 0 else None
    s = str(val).strip()
    if not s:
        return None
    if s.endswith('月'):
        s = s[:-1].strip()
    m = re.match(r'^[\d.]+', s.replace(',', ''))
    if m:
        try:
            n = float(m.group(0))
            return n if n > 0 else None
        except ValueError:
            pass
    m = re.search(r'([\d.]+)\s*months?', s, re.I)
    if m:
        try:
            n = float(m.group(1))
            return n if n > 0 else None
        except ValueError:
            pass
    return None


def _master_status(master_row, matched_no):
    if not master_row:
        return 'missing'
    missing = []
    for db_key, _label, _payload_key in MASTER_COMPARE_FIELDS:
        if db_key in ('quoted_amount', 'awarded_amount'):
            continue
        if not _safe_str(master_row.get(db_key)):
            missing.append(db_key)
    if not _pick_amount(None, master_row):
        missing.append('amount')
    if missing:
        return 'partial'
    return 'ok'


def _merge_item(raw_row, person_code=None, pm_overrides=None):
    pm_overrides = pm_overrides or {}
    master_row, matched_no = find_master_record(raw_row.get('quotation_raw'), person_code)
    quotation_no = (
        matched_no
        or build_quotation_no(raw_row.get('quotation_raw'), person_code)
        or raw_row.get('quotation_raw')
    )
    project_name = (
        _safe_str(master_row.get('description') if master_row else None)
        or raw_row.get('project_name')
    )
    start_date = _safe_date(master_row.get('start_date') if master_row else None) or raw_row.get('start_date')
    completion_date = (
        _safe_date(master_row.get('completion_date') if master_row else None)
        or raw_row.get('completion_date')
    )
    amount = _pick_amount(raw_row.get('amount'), master_row)
    client_name = _infer_client(project_name, master_row)
    person_code_resolved = (
        normalize_person_code(person_code)
        or (master_row.get('person_code') if master_row else None)
        or normalize_person_code(
            (matched_no or '').split('/')[-1] if matched_no else None
        )
    )
    payload = {
        'quotation_raw': raw_row.get('quotation_raw'),
        'quotation_no': quotation_no,
        'core_no': normalize_nn1_quotation(raw_row.get('quotation_raw')),
        'project_name': project_name,
        'client_name': client_name,
        'amount': amount,
        'start_date': start_date,
        'completion_date': completion_date,
        'contract_start_date': (
            pm_overrides.get('contract_start_date')
            or raw_row.get('contract_start_date')
            or None
        ),
        'contract_end_date': (
            pm_overrides.get('contract_end_date')
            or raw_row.get('contract_end_date')
            or None
        ),
        'contract_months': None,
        'expected_period': raw_row.get('expected_period'),
        'subcon_name': (
            raw_row.get('subcon_name')
            or (master_row.get('subcon_company') if master_row else None)
        ),
        'subcon_quotation': raw_row.get('subcon_quotation'),
        'ei_issued': raw_row.get('ei_issued'),
        'email_issued': raw_row.get('email_issued'),
        'person_code': person_code_resolved,
        'person_in_charge': (
            master_row.get('person_in_charge') if master_row else None
        ) or person_display_name(person_code_resolved),
        'approval_category': pm_overrides.get('approval_category') or '合約',
        'submit_date': pm_overrides.get('submit_date'),
        'tender_date': pm_overrides.get('tender_date'),
        'bid_deadline': pm_overrides.get('bid_deadline'),
        'bid_deadline_ampm': pm_overrides.get('bid_deadline_ampm') or '',
        'bid_deadline_meridiem': pm_overrides.get('bid_deadline_meridiem') or '',
        'attachments': pm_overrides.get('attachments') or '',
        'remark': pm_overrides.get('remark') or raw_row.get('remark') or '',
        'master_status': _master_status(master_row, matched_no),
        'master_matched_no': matched_no,
        'master_id': master_row.get('id') if master_row else None,
        'master_gaps': [],
    }
    if master_row:
        gaps = []
        for db_key, label, _pk in MASTER_COMPARE_FIELDS:
            if db_key in ('quoted_amount', 'awarded_amount'):
                if not _pick_amount(raw_row.get('amount'), master_row):
                    gaps.append(label)
                continue
            if not _safe_str(master_row.get(db_key)) and not _safe_str(payload.get(_pk)):
                gaps.append(label)
        payload['master_gaps'] = gaps
    else:
        payload['master_gaps'] = [label for _k, label, _pk in MASTER_COMPARE_FIELDS]
    return payload


def parse_nn1_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_row = None
    col_map = {}
    for r in range(1, min(ws.max_row, 20) + 1):
        row_vals = [_safe_str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        joined = ' '.join(v or '' for v in row_vals)
        if 'Quotation' in joined and ('工程' in joined or '名稱' in joined):
            header_row = r
            for c, val in enumerate(row_vals, 1):
                if not val:
                    continue
                low = val.lower()
                if 'quotation no' in low or low == 'quotation no.':
                    col_map['quotation'] = c
                elif low.startswith('sub-con'):
                    if 'name' in low or 'company' in low or '商' in val:
                        col_map['subcon_name'] = c
                    elif 'subcon' not in col_map:
                        col_map['subcon'] = c
                elif '分判商' in val:
                    col_map['subcon_name'] = c
                elif 'quotation' in low and 'sub-con' not in low and 'quotation' not in col_map:
                    col_map['quotation'] = c
                elif '工程名稱' in val:
                    col_map['name'] = c
                elif '金額' in val:
                    col_map['amount'] = c
                elif '開工' in val:
                    col_map['start'] = c
                elif '完工' in val:
                    col_map['end'] = c
                elif val == 'EI issued':
                    col_map['ei'] = c
                elif val == 'Email issued':
                    col_map['email'] = c
                elif 'Expected Working Period' in val:
                    col_map['period'] = c
            break
    if not header_row or 'quotation' not in col_map:
        raise ValueError('無法辨識 NN1 表頭（需含 Quotation No. 及工程名稱）')

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        raw_q = _safe_str(ws.cell(r, col_map['quotation']).value)
        if not raw_q:
            continue
        name = _safe_str(ws.cell(r, col_map.get('name', 0)).value) if col_map.get('name') else None
        amount = _safe_float(ws.cell(r, col_map.get('amount', 0)).value) if col_map.get('amount') else None
        start = _safe_date(ws.cell(r, col_map.get('start', 0)).value) if col_map.get('start') else None
        end = _safe_date(ws.cell(r, col_map.get('end', 0)).value) if col_map.get('end') else None
        remark = None
        if col_map.get('subcon'):
            tail = _safe_str(ws.cell(r, col_map['subcon'] + 1).value)
            if tail and not re.match(r'^\d', tail):
                remark = tail
        rows.append({
            'quotation_raw': raw_q,
            'project_name': name,
            'amount': amount,
            'start_date': start,
            'completion_date': end,
            'ei_issued': _safe_str(ws.cell(r, col_map.get('ei', 0)).value) if col_map.get('ei') else None,
            'email_issued': _safe_str(ws.cell(r, col_map.get('email', 0)).value) if col_map.get('email') else None,
            'expected_period': _parse_period_months(
                ws.cell(r, col_map.get('period', 0)).value
            ) if col_map.get('period') else None,
            'subcon_name': _safe_str(ws.cell(r, col_map.get('subcon_name', 0)).value) if col_map.get('subcon_name') else None,
            'subcon_quotation': _safe_str(ws.cell(r, col_map.get('subcon', 0)).value) if col_map.get('subcon') else None,
            'remark': remark,
        })
    if not rows:
        raise ValueError('NN1 表內沒有有效項目列')
    return rows


def preview_nn1_import(path, person_code=None):
    raw_rows = parse_nn1_workbook(path)
    items = [_merge_item(row, person_code=person_code) for row in raw_rows]
    stats = {
        'total': len(items),
        'master_ok': sum(1 for i in items if i['master_status'] == 'ok'),
        'master_partial': sum(1 for i in items if i['master_status'] == 'partial'),
        'master_missing': sum(1 for i in items if i['master_status'] == 'missing'),
    }
    return {'items': items, 'stats': stats, 'source_file': os.path.basename(path)}


def rematch_eng_item(item, person_code=None):
    """NN1 編號變更後重新比對 Master（保留 PM／用戶已填欄位）。"""
    item = dict(item or {})
    raw = _safe_str(item.get('quotation_raw'))
    if not raw:
        return item
    pc = person_code or item.get('person_code')
    preserve_keys = (
        'project_name', 'client_name', 'amount',
        'contract_start_date', 'contract_end_date', 'expected_period',
        'subcon_name', 'subcon_quotation',
        'start_date', 'completion_date', 'person_in_charge',
        'approval_category', 'approval_other',
        'submit_date', 'tender_date', 'bid_deadline',
        'bid_deadline_ampm', 'bid_deadline_meridiem',
        'attachments', 'remark',
    )
    preserved = {k: item[k] for k in preserve_keys if item.get(k) not in (None, '')}
    pm_overrides = {
        k: item.get(k) for k in SIGNOFF_PM_FIELDS if item.get(k) not in (None, '')
    }
    raw_row = {
        'quotation_raw': raw,
        'project_name': item.get('project_name'),
        'amount': item.get('amount'),
        'start_date': item.get('start_date'),
        'completion_date': item.get('completion_date'),
        'expected_period': item.get('expected_period'),
        'subcon_name': item.get('subcon_name'),
        'subcon_quotation': item.get('subcon_quotation'),
        'remark': item.get('remark'),
    }
    merged = _merge_item(raw_row, person_code=pc, pm_overrides=pm_overrides)
    merged['quotation_raw'] = raw
    for k, v in preserved.items():
        merged[k] = v
    merged['quotation_no'] = merged.get('master_matched_no') or merged.get('quotation_no')
    return merged


def _coalesce_signoff_date(overrides, base, key):
    """合約年期日期：content 明確傳 null/空字串時視為不填，不 fallback Master。"""
    if key in overrides:
        v = overrides[key]
        return v if v not in (None, '') else None
    v = base.get(key)
    return v if v not in (None, '') else None


def build_signoff_payload(item, pm_overrides=None, content_overrides=None):
    pm_overrides = pm_overrides if pm_overrides is not None else {}
    content_overrides = content_overrides if content_overrides is not None else {}
    base = dict(item)
    person_code = content_overrides.get('person_code') or base.get('person_code')
    contract_start = _coalesce_signoff_date(content_overrides, base, 'contract_start_date')
    contract_end = _coalesce_signoff_date(content_overrides, base, 'contract_end_date')
    payload = _merge_item(
        {
            'quotation_raw': base.get('quotation_raw'),
            'project_name': content_overrides.get('project_name') or base.get('project_name'),
            'amount': content_overrides.get('amount', base.get('amount')),
            'expected_period': content_overrides.get('expected_period') or base.get('expected_period'),
            'contract_start_date': contract_start,
            'contract_end_date': contract_end,
            'subcon_name': content_overrides.get('subcon_name') or base.get('subcon_name'),
            'subcon_quotation': content_overrides.get('subcon_quotation') or base.get('subcon_quotation'),
            'remark': content_overrides.get('remark') or base.get('remark'),
        },
        person_code=person_code,
        pm_overrides={
            **pm_overrides,
            'contract_start_date': contract_start,
            'contract_end_date': contract_end,
        },
    )
    for k in SIGNOFF_CONTENT_FIELDS:
        if k in content_overrides and content_overrides[k] is not None:
            payload[k] = content_overrides[k]
        elif k in ('contract_start_date', 'contract_end_date') and k in content_overrides:
            payload[k] = None
    for k in SIGNOFF_PM_FIELDS:
        if k in pm_overrides:
            payload[k] = pm_overrides[k]
    cat = (payload.get('approval_category') or base.get('approval_category') or '合約').strip()
    payload['approval_category'] = cat if cat in ('合約', '投標意向書', '其他') else '合約'
    if payload.get('amount') not in (None, ''):
        try:
            payload['amount'] = float(str(payload['amount']).replace(',', '').replace('$', ''))
        except (TypeError, ValueError):
            pass
    parsed = _parse_period_months(payload.get('expected_period'))
    if parsed is not None:
        payload['expected_period'] = parsed
    payload['contract_months'] = parsed or _months_between(
        payload.get('contract_start_date'), payload.get('contract_end_date'),
    )
    _sanitize_signoff_contract_period(payload)
    return payload


def _sanitize_signoff_contract_period(payload: dict) -> None:
    """合約年期只用 contract_*；若與 Master 開工/完工相同視為誤填並清空。"""
    cs = payload.get('contract_start_date') or None
    ce = payload.get('contract_end_date') or None
    sd = payload.get('start_date')
    cd = payload.get('completion_date')
    if cs and sd and str(cs)[:10] == str(sd)[:10]:
        if not ce or (cd and str(ce)[:10] == str(cd)[:10]):
            cs = None
            ce = None
    payload['contract_start_date'] = cs
    payload['contract_end_date'] = ce


def resolve_signoff_template(quotation_no=None):
    """統一使用 Ref/投標合約會簽表Template.docx（quotation_no 參數保留相容）。"""
    from signoff_generator import get_signoff_template_path
    return get_signoff_template_path()
