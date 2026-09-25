"""
summary_importer.py — Ref/Summary.xlsx 匯入工程項目主檔

對齊第2頁 Summary：MP合約編號、會計編號、主合約名稱、客方、工程分類、MP承建金額
"""
import os
import re
from io import BytesIO

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import BASE_DIR
from project_cover import derive_mp_contract_code

DEFAULT_SUMMARY_PATH = os.path.join(BASE_DIR, 'Ref', 'Summary.xlsx')

_Q_RE = re.compile(r'^Q\d+_\d+$', re.I)
_N_RE = re.compile(r'^N\d+$', re.I)


def _s(val):
    if val is None:
        return ''
    return str(val).strip()


def _amount(val):
    if val is None or val == '':
        return None
    try:
        n = float(val)
        return n if n != 0 else None
    except (TypeError, ValueError):
        return None


def _is_q_code(val):
    return bool(_Q_RE.match(_s(val)))


def _is_account(val):
    return bool(_N_RE.match(_s(val)))


def _is_main_contractor_line(text):
    t = _s(text)
    return bool(t) and ('主承判商' in t or '主承判' in t)


def _clean_title(title):
    t = _s(title)
    return t or None


def _normalize_client(text):
    """客方：去掉尾斜線、不混入主承判商"""
    t = _s(text).rstrip('/').strip()
    if _is_main_contractor_line(t):
        return None
    return t or None


def _normalize_main_contractor(text):
    """主承判商：去掉 (主承判商) 標記"""
    t = _s(text)
    if not t:
        return None
    t = re.sub(r'\(主承判商\)', '', t)
    t = re.sub(r'\(主承判\)', '', t)
    return t.strip().rstrip('/').strip() or None


def parse_summary_rows(rows):
    """
    解析 Summary 工作表（含 PPT 複製時的拆行、客方續行、金額續行、N23 多 Q 編號）
    回傳 list[dict]
    """
    records = []
    current = None

    def flush():
        nonlocal current
        if not current:
            return
        extra = current.pop('extra_mp_codes', [])
        codes = []
        if current.get('mp_contract_code'):
            codes.append(current['mp_contract_code'])
        for c in extra:
            if c and c not in codes:
                codes.append(c)
        current['mp_contract_codes'] = codes
        if current.get('client'):
            current['client'] = _normalize_client(current['client'])
        if current.get('client_secondary'):
            current['client_secondary'] = _normalize_client(current['client_secondary'])
        if current.get('main_contractor'):
            current['main_contractor'] = _normalize_main_contractor(current['main_contractor'])
        records.append(current)
        current = None

    for row in rows:
        cells = list(row[:6]) + [None] * (6 - len(row))
        mp, acc, title, client, wtype, amt = cells[:6]
        mp_s = _s(mp)
        acc_s = _s(acc)
        title_s = _clean_title(title)
        client_s = _s(client)
        wtype_s = _s(wtype) or None
        amt_v = _amount(amt)

        # 標題列
        if mp_s == 'MP合約編號' or (acc_s == '會計編號' and title_s == '主合約名稱'):
            continue

        # 僅 MP 合約編號續行（N23 等多 Q）
        if _is_q_code(mp_s) and not acc_s and not title_s and not client_s and not amt_v:
            if current:
                current.setdefault('extra_mp_codes', []).append(mp_s.upper())
            continue

        # 僅金額續行 → 補上一条記錄
        if not mp_s and not acc_s and not title_s and not client_s and amt_v:
            target = current
            if not target and records:
                target = records[-1]
            if target and not target.get('contract_amount'):
                target['contract_amount'] = amt_v
            elif current:
                current['contract_amount'] = amt_v
            continue

        # 客方 / 主承建商續行
        if not mp_s and not acc_s and not title_s and client_s and not amt_v:
            if _is_main_contractor_line(client_s):
                if current:
                    mc = _normalize_main_contractor(client_s)
                    if mc:
                        if current.get('main_contractor'):
                            current['main_contractor'] = f"{current['main_contractor']} · {mc}"
                        else:
                            current['main_contractor'] = mc
            elif current:
                part = _normalize_client(client_s)
                if part:
                    if not current.get('client'):
                        current['client'] = part
                    elif not current.get('client_secondary'):
                        current['client_secondary'] = part
                    else:
                        # 第三個及以上客方续并第二客方
                        current['client_secondary'] = (
                            f"{current['client_secondary']} · {part}"
                        )
            continue

        # 新項目列
        if acc_s or title_s or (_is_q_code(mp_s) and (acc_s or title_s or client_s or amt_v)):
            flush()
            current = {
                'mp_contract_code': mp_s.upper() if _is_q_code(mp_s) else None,
                'account_code': acc_s or None,
                'project_name_zh': title_s,
                'client': _normalize_client(client_s),
                'client_secondary': None,
                'work_type': wtype_s,
                'contract_amount': amt_v or 0,
                'main_contractor': None,
                'extra_mp_codes': [],
                'notes': None,
            }
            if not current['mp_contract_code'] and acc_s:
                pass
            continue

    flush()
    return [r for r in records if r.get('account_code') or r.get('project_name_zh')]


def read_summary_file(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return parse_summary_rows(rows)


def sync_summary_import(filepath=None, update_existing=True):
    """匯入 Summary.xlsx → projects 表；回傳統計"""
    import database as db

    path = filepath or DEFAULT_SUMMARY_PATH
    if not os.path.exists(path):
        raise FileNotFoundError(f'找不到 Summary 檔案: {path}')

    rows = read_summary_file(path)
    created = updated = skipped = 0
    details = []

    for row in rows:
        result = db.upsert_summary_project(row, update_existing=update_existing)
        details.append(result)
        if result['action'] == 'created':
            created += 1
        elif result['action'] == 'updated':
            updated += 1
        else:
            skipped += 1

    return {
        'source_file': os.path.basename(path),
        'rows_parsed': len(rows),
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'details': details,
    }


SUMMARY_EXPORT_HEADERS = (
    'MP合約編號', '會計編號', '主合約名稱', '客方', '工程分類', 'MP承建金額',
    '客方二', '主承判商', '負責人', '項目編號', '報價編號', '狀態', '利潤率%', '分判數',
)


def _summary_export_styles():
    fill = PatternFill('solid', fgColor='1F4E79')
    font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    side = Side(style='thin', color='B0B0B0')
    border = Border(left=side, right=side, top=side, bottom=side)
    return fill, font, border


def _summary_autosize(ws, cols):
    for i in range(1, cols + 1):
        letter = get_column_letter(i)
        maxlen = 10
        for cell in ws[letter]:
            v = cell.value
            if v is None:
                continue
            maxlen = max(maxlen, min(len(str(v)), 40))
        ws.column_dimensions[letter].width = maxlen + 2


def _summary_category(row):
    return row.get('category_l2_label') or row.get('category_l2_code') or row.get('work_type')


def _summary_export_rows(items):
    """每項目一行；額外 MP 合約編號以續行輸出（可再匯入）"""
    out = []
    for row in items:
        codes = list(row.get('mp_contract_codes') or [])
        primary = row.get('mp_contract_code')
        if primary and primary not in codes:
            codes.insert(0, primary)
        elif not codes and primary:
            codes = [primary]
        main_mp = codes[0] if codes else None
        out.append([
            main_mp,
            row.get('account_code'),
            row.get('main_contract_title') or row.get('main_contract_title_en'),
            row.get('client'),
            _summary_category(row),
            row.get('contract_amount'),
            row.get('client_secondary'),
            row.get('main_contractor'),
            row.get('person_in_charge'),
            row.get('project_code'),
            row.get('quotation_no'),
            row.get('status'),
            row.get('current_profit_pct'),
            row.get('sc_count'),
        ])
        for extra_mp in codes[1:]:
            out.append([extra_mp, None, None, None, None, None, None, None, None, None, None, None, None, None])
    return out


def export_summary_bytes():
    """匯出工程項目 Summary（對齊 Ref/Summary.xlsx 主欄 + 列表常用欄）"""
    import database as db

    items = db.get_company_summary()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Summary'
    header_fill, header_font, header_border = _summary_export_styles()
    body_font = Font(name='Calibri', size=10)
    for idx, title in enumerate(SUMMARY_EXPORT_HEADERS, start=1):
        cell = ws.cell(1, idx, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
    export_rows = _summary_export_rows(items)
    for r_i, vals in enumerate(export_rows, start=2):
        for c_i, val in enumerate(vals, start=1):
            cell = ws.cell(r_i, c_i, val)
            cell.font = body_font
            cell.border = header_border
    _summary_autosize(ws, len(SUMMARY_EXPORT_HEADERS))
    last_row = max(1, 1 + len(export_rows))
    ws.auto_filter.ref = f'A1:{get_column_letter(len(SUMMARY_EXPORT_HEADERS))}{last_row}'
    ws.freeze_panes = 'B2'
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
