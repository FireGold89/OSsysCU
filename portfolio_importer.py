"""
portfolio_importer.py — N 項目結算總表 / 進度表 Excel 匯入匯出

FA r2：sheet `Final Account`，表頭含 N Code；左 19 欄 + 分判矩陣 15 組。
進度表：11 欄，表頭含 Project Code。
"""
import os
import re
from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import portfolio

FA_SHEET_NAME = 'Final Account'
STATUS_SKIP = frozenset({
    'ON PROGRESS', 'COMPLETED', '待簽', '—', '-', '——', 'NONE', 'NAN',
})
CHECK_YES = frozenset({'✔', '✓', '√', '1', 'Y', 'YES', 'TRUE'})

FA_LEFT_HEADERS = [
    'N Code', 'Project Code', 'Project Description', 'PM',
    'Commencement Date', 'Contract Completion Date', 'PC Date', 'PC Cert',
    'DLP Commencement Date', 'DLP (days)', 'DLP Expiry Date',
    'Retention to be released', 'Defect Correction Certificate',
    '預計完工日期', 'Remark', 'Contract Sum',
    'Project Completed\n/On Progress', 'Client', 'Client \nFinal Account Status',
]
PROGRESS_HEADERS = [
    'Project Code', 'Project Description', 'Client', 'PM',
    'Commencement Date', 'Contract Completion Date', 'PC Date',
    '預計完工日期', 'Remark', 'Contract Sum', 'Completed\n/On Progress',
]


def _safe_str(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if not s or s.lower() in ('none', 'nan'):
        return None
    if s.upper() in ('NA', 'N/A', '-', '—', '——'):
        return None
    return s


def _safe_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    s = _safe_str(val)
    if not s:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d/%m/%y', '%Y/%m/%d'):
        try:
            return datetime.strptime(s[:10], fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    if re.match(r'\d{4}-\d{2}-\d{2}', s):
        return s[:10]
    return None


def _safe_float(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if isinstance(val, float) and val != val:
            return None
        return float(val)
    s = _safe_str(val)
    if not s or s == '?':
        return None
    s = s.replace(',', '').replace('$', '').replace('HK', '')
    try:
        return float(s)
    except ValueError:
        return None


def _safe_int(val):
    n = _safe_float(val)
    if n is None:
        return None
    return int(n)


def _is_checked(val):
    if val is True or val == 1:
        return 1
    s = _safe_str(val)
    if not s:
        return 0
    return 1 if s.upper() in CHECK_YES or s == '✔' else 0


def _norm_status(val, allowed):
    s = _safe_str(val)
    if not s:
        return None
    s = re.sub(r'\s+', ' ', s).strip()
    lower = s.lower()
    if lower == 'on progress':
        s = 'On Progress'
    elif lower == 'completed':
        s = 'Completed'
    if s in allowed:
        return s
    return None


def _looks_like_project_code(code):
    s = (code or '').strip()
    if not s or s.upper() in STATUS_SKIP:
        return False
    if re.search(r'(?:MS/)?Q\d', s, re.I):
        return True
    if re.match(r'^Q\d', s, re.I):
        return True
    return False


def _cell(ws, row, col):
    return ws.cell(row, col).value


def _find_header_row(ws, needle, max_scan=8):
    target = needle.lower()
    for r in range(1, min(max_scan, ws.max_row) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = _cell(ws, r, c)
            if v and str(v).strip().lower() == target:
                return r
    return None


def parse_fa_list(path):
    """解析 FA r2 Excel → list of dict。略過狀態下拉列。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == FA_SHEET_NAME.lower():
            sheet = name
            break
    if not sheet:
        sheet = wb.sheetnames[0]
    ws = wb[sheet]
    header_row = _find_header_row(ws, 'N Code')
    if not header_row:
        raise ValueError('找不到 N Code 表頭（請上傳 Final Account Status List r2）')
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        n_code = _safe_str(_cell(ws, r, 1))
        project_code = _safe_str(_cell(ws, r, 2))
        if not _looks_like_project_code(project_code):
            continue
        slots = []
        for i in range(15):
            name_col = 20 + i * 2
            st_col = name_col + 1
            name = _safe_str(_cell(ws, r, name_col))
            st = _norm_status(_cell(ws, r, st_col), portfolio.CLIENT_FAC_STATUS_VALUES)
            if not name and not st:
                continue
            slots.append({'slot': i + 1, 'name': name or '', 'fac_status': st or ''})
        rows.append({
            'n_code': n_code,
            'project_code': project_code,
            'description': _safe_str(_cell(ws, r, 3)),
            'pm': _safe_str(_cell(ws, r, 4)),
            'commencement_date': _safe_date(_cell(ws, r, 5)),
            'contract_completion_date': _safe_date(_cell(ws, r, 6)),
            'pc_date': _safe_date(_cell(ws, r, 7)),
            'pc_cert_done': _is_checked(_cell(ws, r, 8)),
            'dlp_commencement_date': _safe_date(_cell(ws, r, 9)),
            'dlp_days': _safe_int(_cell(ws, r, 10)),
            'dlp_expiry_date': _safe_date(_cell(ws, r, 11)),
            'retention_to_release': _safe_float(_cell(ws, r, 12)),
            'defect_cert_done': _is_checked(_cell(ws, r, 13)),
            'expected_completion_date': _safe_date(_cell(ws, r, 14)),
            'remark': _safe_str(_cell(ws, r, 15)) or '',
            'contract_sum': _safe_float(_cell(ws, r, 16)),
            'project_progress_status': _norm_status(
                _cell(ws, r, 17), portfolio.PROGRESS_STATUS_VALUES,
            ) or 'On Progress',
            'client': _safe_str(_cell(ws, r, 18)),
            'client_fac_status': _norm_status(
                _cell(ws, r, 19), portfolio.CLIENT_FAC_STATUS_VALUES,
            ) or 'On Progress',
            'subcontractors': slots,
        })
    return {'sheet': sheet, 'header_row': header_row, 'rows': rows}


def parse_progress_list(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_row = _find_header_row(ws, 'Project Code')
    if not header_row:
        raise ValueError('找不到 Project Code 表頭（請上傳 On Progress Projects）')
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        project_code = _safe_str(_cell(ws, r, 1))
        if not _looks_like_project_code(project_code):
            continue
        rows.append({
            'project_code': project_code,
            'description': _safe_str(_cell(ws, r, 2)),
            'client': _safe_str(_cell(ws, r, 3)),
            'pm': _safe_str(_cell(ws, r, 4)),
            'commencement_date': _safe_date(_cell(ws, r, 5)),
            'contract_completion_date': _safe_date(_cell(ws, r, 6)),
            'pc_date': _safe_date(_cell(ws, r, 7)),
            'expected_completion_date': _safe_date(_cell(ws, r, 8)),
            'remark': _safe_str(_cell(ws, r, 9)) or '',
            'contract_sum': _safe_float(_cell(ws, r, 10)),
            'project_progress_status': _norm_status(
                _cell(ws, r, 11), portfolio.PROGRESS_STATUS_VALUES,
            ) or 'On Progress',
            'n_code': None,
            'client_fac_status': None,
            'pc_cert_done': None,
            'defect_cert_done': None,
            'subcontractors': [],
        })
    return {'sheet': ws.title, 'header_row': header_row, 'rows': rows}


def preview_fa_list(path):
    parsed = parse_fa_list(path)
    return _preview_rows(parsed, os.path.basename(path), 'fa_list')


def preview_progress_list(path):
    parsed = parse_progress_list(path)
    return _preview_rows(parsed, os.path.basename(path), 'progress_list')


def _preview_rows(parsed, filename, import_type):
    matched = []
    unmatched = []
    for row in parsed['rows']:
        pid = portfolio.find_project_id_for_code(row['project_code'])
        item = {
            'n_code': row.get('n_code') or '',
            'project_code': row['project_code'],
            'description': row.get('description') or '',
            'status': row.get('project_progress_status'),
        }
        if pid:
            item['project_id'] = pid
            matched.append(item)
        else:
            unmatched.append(item)
    return {
        'import_type': import_type,
        'filename': filename,
        'sheet': parsed.get('sheet'),
        'rows_read': len(parsed['rows']),
        'matched': len(matched),
        'unmatched': len(unmatched),
        'matched_rows': matched[:80],
        'unmatched_rows': unmatched[:80],
        'create_placeholders': True,
    }


def sync_fa_list(path, create_placeholders=True):
    parsed = parse_fa_list(path)
    return _sync_rows(
        parsed, os.path.basename(path), 'fa_list',
        create_placeholders=create_placeholders, write_slots=True,
    )


def sync_progress_list(path, create_placeholders=True):
    parsed = parse_progress_list(path)
    return _sync_rows(
        parsed, os.path.basename(path), 'progress_list',
        create_placeholders=create_placeholders, write_slots=False,
    )


def _sync_rows(parsed, filename, import_type, create_placeholders, write_slots):
    upserted = 0
    created_projects = 0
    skipped = 0
    slots_written = 0
    for row in parsed['rows']:
        result = portfolio.upsert_from_import(
            row,
            create_placeholder=create_placeholders,
            write_slots=write_slots,
            source_file=filename,
        )
        if result.get('action') == 'skipped':
            skipped += 1
            continue
        upserted += 1
        if result.get('created_project'):
            created_projects += 1
        slots_written += int(result.get('slots') or 0)
    rec = portfolio.record_import(
        import_type, filename, len(parsed['rows']), upserted,
    )
    return {
        'import_type': import_type,
        'filename': filename,
        'rows_read': len(parsed['rows']),
        'rows_upserted': upserted,
        'created_projects': created_projects,
        'skipped': skipped,
        'slots_written': slots_written,
        'imported_at': rec.get('imported_at'),
    }


def _excel_date(val):
    s = _safe_str(val)
    if not s:
        return None
    try:
        return datetime.strptime(s[:10], '%Y-%m-%d').date()
    except ValueError:
        return s


def _excel_check(done):
    return '✔' if done else None


def _header_font():
    return Font(name='Calibri', bold=True, size=10)


def _thin_border():
    side = Side(style='thin', color='B0B0B0')
    return Border(left=side, right=side, top=side, bottom=side)


def export_fa_list_bytes():
    data = portfolio.list_portfolio(status='all', include_subcons=True)
    items = data.get('items') or []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = FA_SHEET_NAME
    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='center')
    for i in range(15):
        col = 20 + i * 2
        cell = ws.cell(1, col, f'Sub-Contractor {i + 1}')
        cell.font = header_font
        cell.fill = header_fill
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
    headers = list(FA_LEFT_HEADERS)
    for _i in range(15):
        headers.append('Name of Sub-contractor')
        headers.append('Final Account Status')
    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(2, idx, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        cell.border = _thin_border()
    for r_i, item in enumerate(items, start=3):
        vals = [
            item.get('n_code') or None,
            item.get('project_code'),
            item.get('description'),
            item.get('pm') or None,
            _excel_date(item.get('commencement_date')),
            _excel_date(item.get('contract_completion_date')),
            _excel_date(item.get('pc_date')),
            _excel_check(item.get('pc_cert_done')),
            _excel_date(item.get('dlp_commencement_date')),
            item.get('dlp_days'),
            _excel_date(item.get('dlp_expiry_date')),
            item.get('retention_to_release'),
            _excel_check(item.get('defect_cert_done')),
            _excel_date(item.get('expected_completion_date')),
            item.get('remark') or None,
            item.get('contract_sum'),
            item.get('project_progress_status'),
            item.get('client') or None,
            item.get('client_fac_status') or None,
        ]
        slot_map = {s.get('slot'): s for s in (item.get('subcontractors') or [])}
        for i in range(1, 16):
            sl = slot_map.get(i) or {}
            vals.append(sl.get('name') or None)
            vals.append(sl.get('fac_status') or None)
        for c_i, val in enumerate(vals, start=1):
            cell = ws.cell(r_i, c_i, val)
            cell.border = _thin_border()
            cell.font = Font(name='Calibri', size=10)
    _autosize(ws, 49)
    ws.auto_filter.ref = f'A2:{get_column_letter(49)}{max(2, 2 + len(items))}'
    ws.freeze_panes = 'D3'
    ws.row_dimensions[2].height = 32
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_progress_list_bytes():
    data = portfolio.list_portfolio(status='all', include_subcons=False)
    items = data.get('items') or []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'On Progress Projects'
    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='center')
    for idx, title in enumerate(PROGRESS_HEADERS, start=1):
        cell = ws.cell(1, idx, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        cell.border = _thin_border()
    for r_i, item in enumerate(items, start=2):
        vals = [
            item.get('project_code'),
            item.get('description'),
            item.get('client') or None,
            item.get('pm') or None,
            _excel_date(item.get('commencement_date')),
            _excel_date(item.get('contract_completion_date')),
            _excel_date(item.get('pc_date')),
            _excel_date(item.get('expected_completion_date')),
            item.get('remark') or None,
            item.get('contract_sum'),
            item.get('project_progress_status'),
        ]
        for c_i, val in enumerate(vals, start=1):
            cell = ws.cell(r_i, c_i, val)
            cell.border = _thin_border()
            cell.font = Font(name='Calibri', size=10)
    _autosize(ws, 11)
    ws.auto_filter.ref = f'A1:K{max(1, 1 + len(items))}'
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 28
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _autosize(ws, cols):
    for i in range(1, cols + 1):
        letter = get_column_letter(i)
        maxlen = 10
        for cell in ws[letter]:
            v = cell.value
            if v is None:
                continue
            maxlen = max(maxlen, min(len(str(v)), 36))
        ws.column_dimensions[letter].width = maxlen + 2
